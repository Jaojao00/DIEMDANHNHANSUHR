from docx import Document
from docx.shared import Pt
from docx.shared import RGBColor

doc = Document()

# Title
title = doc.add_heading('TÀI LIỆU CHI TIẾT HỆ THỐNG ĐIỂM DANH & QUẢN LÝ NHÂN SỰ AGARI', 0)
title.style.font.color.rgb = RGBColor(230, 60, 30)

# Section 1
doc.add_heading('1. TỔNG QUAN HỆ THỐNG', level=1)
doc.add_paragraph('Hệ thống phần mềm nội bộ (Web Application) quản lý điểm danh và tra cứu công lương dành riêng cho đội ngũ nhân sự và ban quản lý. Hệ thống hoạt động theo kiến trúc Client-Server (Frontend + Backend Google Apps Script) không cần server lưu trữ riêng.')
doc.add_paragraph('Tự động hóa: Giảm 80% thời gian xử lý thủ công của HR.', style='List Bullet')
doc.add_paragraph('Minh bạch: Cung cấp công cụ tra cứu công/lương/KPI tức thời cho nhân viên.', style='List Bullet')
doc.add_paragraph('Realtime: Trang bị hệ thống quản trị (Admin Dashboard) cập nhật trạng thái nhân viên theo thời gian thực.', style='List Bullet')

# Section 2
doc.add_heading('2. CHI TIẾT CÁC TÍNH NĂNG (FEATURES)', level=1)

doc.add_heading('A. Giao Diện Nhân Viên (Employee Portal)', level=2)
doc.add_paragraph('1. Điểm Danh: Điểm danh đầu ca, cuối ca, xác nhận vị trí làm việc. Hiển thị thông báo trạng thái realtime bằng Modal và Toast.', style='List Bullet')
doc.add_paragraph('2. Lịch Làm Việc: Xem danh sách các ca làm việc đang mở. Đăng ký ca làm, đăng ký vị trí.', style='List Bullet')
doc.add_paragraph('3. Xem Lịch: Xem chi tiết danh sách những người làm cùng ca, kiểm tra vị trí được phân công.', style='List Bullet')
doc.add_paragraph('4. Tra Cứu Lương & KPI: Tính năng mới giúp đối soát KPI, xem tổng công, số tiền tạm tính, biểu đồ số giờ làm việc. Tích hợp nhập CCCD/Mã NV để bảo mật. Tự động hiển thị biểu đồ Bar Chart thống kê công từng ngày.', style='List Bullet')
doc.add_paragraph('5. Quy Định: Bảng nội quy công ty, quy chế làm việc, thưởng phạt được tích hợp trực tiếp giúp nhân viên dễ dàng nắm bắt.', style='List Bullet')

doc.add_heading('B. Giao Diện Quản Trị (Admin Dashboard)', level=2)
doc.add_paragraph('1. Realtime Tracking: Quản lý nhìn thấy tức thời trạng thái điểm danh của từng nhân sự, số lượng hiện diện/vắng mặt. Bảng điều khiển (Dashboard) với các thông số tổng quan (Tổng NV, Đã Điểm Danh, Chưa Điểm Danh, Xin OFF).', style='List Bullet')
doc.add_paragraph('2. Quản Lý Yêu Cầu (Change Requests): Duyệt hoặc từ chối các yêu cầu xin OFF, đổi ca, đổi áo, đổi thẻ.', style='List Bullet')
doc.add_paragraph('3. Duyệt Điểm Danh & Vị Trí: Quản lý trực tiếp phân vị trí làm việc hàng loạt, điều chỉnh trạng thái điểm danh thủ công.', style='List Bullet')
doc.add_paragraph('4. Đồng Bộ (Sync): Hệ thống tự động đồng bộ (Auto Save) toàn bộ dữ liệu trạng thái xuống Google Sheets thông minh, đảm bảo không ghi đè dữ liệu điểm danh đã có.', style='List Bullet')

# Section 3
doc.add_heading('3. GIAO DIỆN & THIẾT KẾ (UI/UX - TONE MÀU & CHỮ)', level=1)
doc.add_paragraph('Hệ thống sử dụng giao diện Dark Mode chủ đạo, lấy cảm hứng từ phong cách hiện đại, kết hợp Theme "Đêm Hội Trăng Rằm" (Mid-Autumn) sang trọng.')

doc.add_heading('3.1. Bảng màu (Color Palette):', level=2)
doc.add_paragraph('Nền chính (Background): Sử dụng dải màu từ #080A18 (Xanh tím than đậm) đến #1A1A2E (Tím than nhạt), tạo chiều sâu.', style='List Bullet')
doc.add_paragraph('Bề mặt (Surface/Glassmorphism): rgba(20, 20, 40, 0.80) kết hợp hiệu ứng Blur tạo cảm giác kính mờ sang trọng.', style='List Bullet')
doc.add_paragraph('Màu chủ đạo (Primary Brand): #E63C1E (Đỏ Agari) và #FF8C42 (Cam sáng) dùng cho nút bấm, logo, thanh loading.', style='List Bullet')
doc.add_paragraph('Màu trạng thái (Status): Thành công #00D48E (Xanh ngọc), Cảnh báo #FFB800 (Vàng), Lỗi/Xóa #FF4757 (Đỏ nhạt).', style='List Bullet')
doc.add_paragraph('Màu chữ (Text): #F0F0FF (Trắng sáng) cho tiêu đề chính, #9898C0 (Xám tím) cho chữ phụ/mô tả.', style='List Bullet')

doc.add_heading('3.2. Phông chữ (Typography):', level=2)
doc.add_paragraph('Phông chữ chính (Inter): Là phông chữ không chân (sans-serif) hiện đại, thanh thoát, dễ đọc. Được dùng cho văn bản chính, tên nút bấm, giao diện đọc quy định.', style='List Bullet')
doc.add_paragraph('Phông chữ hệ thống/số liệu (JetBrains Mono): Là phông chữ Monospace chuyên dụng, tạo cảm giác công nghệ, chính xác. Được dùng đặc biệt cho các con số, mã nhân viên (VD: OPS12345), và giờ làm việc (VD: 06:00 - 15:00).', style='List Bullet')

# Section 4
doc.add_heading('4. KIẾN TRÚC & CÔNG NGHỆ ÁP DỤNG', level=1)
doc.add_paragraph('Giao diện (Frontend): Viết bằng thuần HTML5, CSS3, Vanilla JavaScript (Không sử dụng Framework cồng kềnh), giúp ứng dụng có tốc độ tải siêu tốc. Tích hợp thư viện SweetAlert2 cho thông báo và Chart.js cho biểu đồ thống kê.', style='List Bullet')
doc.add_paragraph('Trải nghiệm đa nền tảng: Thiết kế Mobile-first responsive chuẩn mobile, tích hợp kỹ thuật CSS Breakout để tự động mở rộng không gian tối đa (100vw) trên màn hình máy tính (PC).', style='List Bullet')
doc.add_paragraph('PWA (Progressive Web App): Ứng dụng tích hợp Service Worker (sw.js) để lưu cache, cho phép load cực nhanh cả khi rớt mạng, hỗ trợ lưu trữ Offline và cho phép cài đặt trực tiếp vào màn hình chính của điện thoại như App Native.', style='List Bullet')
doc.add_paragraph('Máy chủ (Backend): Hệ thống sử dụng Google Apps Script (GAS) làm REST API, kết nối với cơ sở dữ liệu là Google Sheets. Giải pháp này tiết kiệm 100% chi phí duy trì server, độ ổn định cực cao.', style='List Bullet')

doc.save('TaiLieuHeThong.docx')
